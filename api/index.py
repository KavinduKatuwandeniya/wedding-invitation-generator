import io
import os
import re
import zipfile

import fitz
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook

FONT_SIZE = 10
LETTER_SPACING = 0.84
FONT_COLOR = (118 / 255, 88 / 255, 48 / 255)

LINE_CENTER_X = 199.001
BASELINE_Y = 134.0
DOTTED_LINE_Y = 138.5

LINE_LEFT_X = 99.0
LINE_RIGHT_X = 299.0
DOTTED_LEFT_X = 99.57
DOTTED_RIGHT_X = 298.43

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROCKWELL_FONT = os.path.join(BASE_DIR, "fonts", "ROCK.TTF")

app = FastAPI()


def sanitize_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]', "_", name)
    return name.strip()[:150]


def normalize_name(value: str) -> str:
    return str(value).strip().upper()


def read_names_from_excel(excel_bytes: bytes) -> list[str]:
    workbook = load_workbook(
        filename=io.BytesIO(excel_bytes),
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active
    names = []

    for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=1):
        value = row[0].value

        if value is None:
            continue

        name = normalize_name(value)

        if name:
            names.append(name)

    workbook.close()
    return names


def create_invitation(template_bytes: bytes, name: str) -> bytes:
    document = fitz.open(stream=template_bytes, filetype="pdf")
    page = document[0]

    page.insert_font(
        fontname="RockwellExact",
        fontfile=ROCKWELL_FONT,
    )

    font = fitz.Font(fontfile=ROCKWELL_FONT)

    # Remove the original dotted-line area so the generated name
    # and replacement dotted line can be positioned precisely.
    page.draw_rect(
        fitz.Rect(
            LINE_LEFT_X,
            132.0,
            LINE_RIGHT_X,
            140.0,
        ),
        color=None,
        fill=(1, 1, 1),
        overlay=True,
    )

    normal_text_width = font.text_length(
        name,
        fontsize=FONT_SIZE,
    )

    spacing_width = LETTER_SPACING * max(len(name) - 1, 0)
    total_text_width = normal_text_width + spacing_width

    current_x = LINE_CENTER_X - (total_text_width / 2)

    for character in name:
        page.insert_text(
            point=(current_x, BASELINE_Y),
            text=character,
            fontsize=FONT_SIZE,
            fontname="RockwellExact",
            color=FONT_COLOR,
            overlay=True,
        )

        character_width = font.text_length(
            character,
            fontsize=FONT_SIZE,
        )

        current_x += character_width + LETTER_SPACING

    # Draw the dotted line below the name.
    page.draw_line(
        p1=(DOTTED_LEFT_X, DOTTED_LINE_Y),
        p2=(DOTTED_RIGHT_X, DOTTED_LINE_Y),
        color=FONT_COLOR,
        width=0.65,
        dashes="[0.5 2.0] 0",
        overlay=True,
    )

    output = io.BytesIO()

    document.save(
        output,
        garbage=4,
        deflate=True,
    )

    document.close()

    return output.getvalue()


def validate_template(pdf_bytes: bytes):
    try:
        document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf",
        )

        if document.page_count < 1:
            raise ValueError("The PDF does not contain any pages.")

        document.close()
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid PDF template: {exc}",
        )


@app.post("/api/generate")
async def generate_invitations(
    pdf: UploadFile = File(...),
    name: str | None = Form(None),
    excel: UploadFile | None = File(None),
):
    if not os.path.exists(ROCKWELL_FONT):
        raise HTTPException(
            status_code=500,
            detail="Rockwell font file is missing from the server.",
        )

    if not pdf.filename or not pdf.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="Please upload a PDF invitation template.",
        )

    has_name = bool(name and name.strip())
    has_excel = excel is not None and bool(excel.filename)

    if not has_name and not has_excel:
        raise HTTPException(
            status_code=400,
            detail="Enter a name or upload an Excel file.",
        )

    if has_name and has_excel:
        raise HTTPException(
            status_code=400,
            detail="Use either a name or an Excel file, not both.",
        )

    pdf_bytes = await pdf.read()
    validate_template(pdf_bytes)

    # Single-name mode
    if has_name:
        normalized_name = normalize_name(name)

        if not normalized_name:
            raise HTTPException(
                status_code=400,
                detail="Please enter a name.",
            )

        try:
            pdf_output = create_invitation(
                pdf_bytes,
                normalized_name,
            )
        except Exception as exc:
            raise HTTPException(
                status_code=500,
                detail=f"Failed to generate invitation: {exc}",
            )

        filename = (
            f"Invitation_{sanitize_filename(normalized_name)}.pdf"
        )

        return Response(
            content=pdf_output,
            media_type="application/pdf",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="{filename}"'
                )
            },
        )

    # Excel bulk mode
    if not excel.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=400,
            detail="Please upload an Excel file (.xlsx or .xlsm).",
        )

    excel_bytes = await excel.read()

    try:
        names = read_names_from_excel(excel_bytes)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read Excel file: {exc}",
        )

    if not names:
        raise HTTPException(
            status_code=400,
            detail="No names were found. Names should start from cell A2.",
        )

    zip_buffer = io.BytesIO()

    try:
        with zipfile.ZipFile(
            zip_buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as zip_file:
            for index, guest_name in enumerate(names, start=1):
                pdf_output = create_invitation(
                    pdf_bytes,
                    guest_name,
                )

                filename = (
                    f"{index:03d}_"
                    f"Invitation_"
                    f"{sanitize_filename(guest_name)}.pdf"
                )

                zip_file.writestr(
                    filename,
                    pdf_output,
                )

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate invitations: {exc}",
        )

    return Response(
        content=zip_buffer.getvalue(),
        media_type="application/zip",
        headers={
            "Content-Disposition": (
                'attachment; filename="Wedding_Invitations.zip"'
            )
        },
    )
